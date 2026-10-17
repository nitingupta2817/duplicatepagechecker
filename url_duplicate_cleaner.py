"""Run with: python -m streamlit run url_duplicate_cleaner.py"""
import csv
import hashlib
from collections import Counter
from copy import copy
from io import BytesIO, StringIO
from urllib.parse import urlparse

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


def parse_paste(text):
    """Excel clipboard is TSV; preserve empty cells and quoted multiline cells."""
    rows = list(csv.reader(StringIO(text), delimiter="\t"))
    if not rows:
        raise ValueError("Paste some Excel cells first.")
    width = max(map(len, rows))
    if width == 0:
        raise ValueError("No cells found.")
    return [r + [""] * (width - len(r)) for r in rows]


def inspect_duplicates(grid, axis, position, skip, trim=True):
    records = grid if axis == "Rows" else list(map(list, zip(*grid)))
    keys = [str(r[position] if r[position] is not None else "") for r in records]
    if trim:
        keys = [k.strip() for k in keys]
    keys = ["" if i < skip else k for i, k in enumerate(keys)]
    counts = Counter(k for k in keys if k)
    seen = set()
    status = []
    for k in keys:
        if not k:
            status.append("Blank / excluded")
        elif counts[k] == 1:
            status.append("Unique")
        elif k in seen:
            status.append("Duplicate")
        else:
            status.append("First occurrence")
        seen.add(k)
    return records, keys, status


def export_excel(grid, removed, axis, source=None, sheet=None):
    if source:
        if not removed:
            return source
        wb = load_workbook(BytesIO(source))
        ws = wb[sheet]
        # Deleting cells does not safely update these dependent Excel features.
        if (any(c.data_type == "f" for s in wb for row in s for c in row)
                or any(s.tables or s.defined_names for s in wb)
                or wb.defined_names or ws.merged_cells.ranges
                or ws.conditional_formatting or ws.data_validations.count
                or any(s._charts or s._images for s in wb)):
            raise ValueError("This workbook has formulas, tables, merged cells, named ranges, "
                             "charts, images or rules that may break when cells shift. "
                             "Copy the plain data and use Paste from Excel instead.")
        dimensions = (dict(ws.row_dimensions) if axis == "Rows"
                      else dict(ws.column_dimensions))
        if axis == "Columns" and any(d.min != d.max for d in dimensions.values()):
            raise ValueError("Grouped column formatting is unsupported. Use Paste from Excel.")
        for i in sorted(removed, reverse=True):
            if axis == "Rows":
                ws.delete_rows(i + 1)
            else:
                ws.delete_cols(i + 1)
        target = ws.row_dimensions if axis == "Rows" else ws.column_dimensions
        target.clear()
        limit = len(grid) if axis == "Rows" else len(grid[0])
        new_i = 0
        for old_i in range(limit):
            if old_i in removed:
                continue
            old_key = old_i + 1 if axis == "Rows" else get_column_letter(old_i + 1)
            new_key = new_i + 1 if axis == "Rows" else get_column_letter(new_i + 1)
            if old_key in dimensions:
                d = copy(dimensions[old_key])
                d.index = new_key
                if axis == "Columns":
                    d.min = d.max = new_i + 1
                target[new_key] = d
            new_i += 1
        if ws.auto_filter.ref:
            ws.auto_filter.ref = ws.calculate_dimension()
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Cleaned Data"
        remaining = ([r for i, r in enumerate(grid) if i not in removed]
                     if axis == "Rows" else
                     [[v for i, v in enumerate(r) if i not in removed] for r in grid])
        for row in remaining:
            ws.append(row)
        # Clipboard strings must remain literal, including values starting with '='.
        for row in ws:
            for c in row:
                if isinstance(c.value, str):
                    c.data_type = "s"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


# ---------------------------------------------------------------------------
# URL Analyzer page helpers
# ---------------------------------------------------------------------------

def parse_url_list(text, trim=True):
    """One URL per cell: splits on new lines and tabs (works for an Excel column or row)."""
    cells = [c for line in text.splitlines() for c in line.split("\t")]
    if trim:
        cells = [c.strip() for c in cells]
    return [c for c in cells if c.strip()]


def extract_domain(url, merge_www=True):
    """Return the lower-case host name, e.g. 'https://WWW.Example.com:8080/a' -> 'example.com'."""
    u = url.strip()
    try:
        host = urlparse(u if "://" in u else "//" + u).hostname or ""
    except ValueError:
        host = ""
    host = host.lower().rstrip(".")
    if merge_www and host.startswith("www."):
        host = host[4:]
    return host


def analyze_urls(urls, merge_www=True):
    url_counts = Counter(urls)
    domains = [extract_domain(u, merge_www) or "(no domain found)" for u in urls]

    per_domain = {}
    for url, dom in zip(urls, domains):
        d = per_domain.setdefault(dom, {"entries": 0, "urls": set()})
        d["entries"] += 1
        d["urls"].add(url)
    domain_table = pd.DataFrame(
        [{"Domain": dom,
          "Total URLs": d["entries"],
          "Different URLs": len(d["urls"]),
          "Exact duplicate copies": d["entries"] - len(d["urls"])}
         for dom, d in per_domain.items()]
    ).sort_values(["Total URLs", "Domain"], ascending=[False, True], ignore_index=True)

    dup_table = pd.DataFrame(
        [{"URL": u, "Domain": extract_domain(u, merge_www) or "(no domain found)",
          "Times found": n, "Extra copies": n - 1}
         for u, n in url_counts.items() if n > 1],
        columns=["URL", "Domain", "Times found", "Extra copies"],
    ).sort_values(["Times found", "URL"], ascending=[False, True], ignore_index=True)

    domain_counts = Counter(domains)
    all_table = pd.DataFrame({
        "Position": range(1, len(urls) + 1),
        "URL": urls,
        "Domain": domains,
        "URL count": [url_counts[u] for u in urls],
        "URLs on same domain": [domain_counts[d] for d in domains],
    })

    summary = {
        "Total URLs pasted": len(urls),
        "Different (unique) URLs": len(url_counts),
        "Exact duplicate copies": len(urls) - len(url_counts),
        "URLs that are repeated": len(dup_table),
        "Different domains": len(per_domain),
        "Domains with more than one URL": int((domain_table["Total URLs"] > 1).sum()),
        "URLs sharing a domain with another URL": int(
            domain_table.loc[domain_table["Total URLs"] > 1, "Total URLs"].sum()),
    }
    return summary, domain_table, dup_table, all_table


def report_excel(summary, domain_table, dup_table, all_table):
    wb = Workbook()
    sheets = [("Summary", pd.DataFrame(list(summary.items()), columns=["Metric", "Value"])),
              ("By domain", domain_table), ("Exact duplicates", dup_table), ("All URLs", all_table)]
    for n, (title, df) in enumerate(sheets):
        ws = wb.active if n == 0 else wb.create_sheet()
        ws.title = title
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([v.item() if hasattr(v, "item") else v for v in row])
        for row in ws:
            for c in row:
                if isinstance(c.value, str):
                    c.data_type = "s"  # keep text literal, even if it starts with '='
        for i, col in enumerate(df.columns, 1):
            width = max([len(str(col))] + [len(str(v)) for v in df[col].head(500)])
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 80)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def analyzer_page():
    st.title("URL Analyzer")
    st.write("Paste a bulk list of URLs to see how many are different, how many share a domain, "
             "and how many are exact duplicates.")
    text = st.text_area("Paste URLs (one per line, or copy an Excel column / row)", height=250,
                        placeholder="https://example.com/a\nhttps://example.com/b\nhttps://example.com/a\nhttps://another.org/page")
    c1, c2 = st.columns(2)
    trim = c1.checkbox("Ignore spaces before / after URLs", value=True, key="an_trim")
    merge_www = c2.checkbox("Treat www.example.com and example.com as the same domain", value=True)
    st.caption("Exact duplicate = the full URL text matches exactly (case, http/https, trailing "
               "slashes and query parameters count as different). Blank cells are ignored.")
    if not text.strip():
        return
    urls = parse_url_list(text, trim)
    if not urls:
        st.warning("No URLs found in the pasted text.")
        return

    summary, domain_table, dup_table, all_table = analyze_urls(urls, merge_www)

    st.subheader("Summary")
    m = st.columns(4)
    m[0].metric("Total URLs pasted", summary["Total URLs pasted"])
    m[1].metric("Different URLs", summary["Different (unique) URLs"])
    m[2].metric("Exact duplicate copies", summary["Exact duplicate copies"],
                help="Extra copies beyond the first of each URL. Deleting these leaves only different URLs.")
    m[3].metric("URLs that are repeated", summary["URLs that are repeated"],
                help="How many distinct URLs appear more than once.")
    m = st.columns(4)
    m[0].metric("Different domains", summary["Different domains"])
    m[1].metric("Domains with 2+ URLs", summary["Domains with more than one URL"])
    m[2].metric("URLs sharing a domain", summary["URLs sharing a domain with another URL"],
                help="URLs whose domain also appears on at least one other pasted URL.")

    tab1, tab2, tab3 = st.tabs(["Same domain", "Exact duplicates", "All URLs"])
    with tab1:
        st.caption("Each domain with how many URLs point to it.")
        st.dataframe(domain_table, hide_index=True, use_container_width=True)
    with tab2:
        if dup_table.empty:
            st.success("No exact duplicate URLs found.")
        else:
            st.caption("URLs that appear more than once in your list.")
            st.dataframe(dup_table, hide_index=True, use_container_width=True)
    with tab3:
        def highlight(row):
            color = "#ffe0df" if row["URL count"] > 1 else ""
            return [f"background-color: {color}; color: #222" if color else ""] * len(row)
        st.caption("Red = this exact URL appears more than once.")
        st.dataframe(all_table.style.apply(highlight, axis=1), hide_index=True,
                     use_container_width=True)

    st.download_button("Download report (Excel)",
                       report_excel(summary, domain_table, dup_table, all_table),
                       "url_analysis.xlsx",
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------------------------------------------------------------------------
# Duplicate Cleaner page (original tool)
# ---------------------------------------------------------------------------

def cleaner_page():
    st.title("URL Duplicate Cleaner")
    st.write("Paste Excel cells, review highlighted duplicate URLs, then delete in bulk and download.")
    mode = st.radio("Input", ["Paste from Excel", "Upload Excel (.xlsx)"], horizontal=True)
    source = None
    sheet = None
    try:
        if mode == "Paste from Excel":
            text = st.text_area("Copy your Excel range and paste here", height=200,
                                placeholder="URL\tTitle\nhttps://example.com/a\tArticle A\nhttps://example.com/a\tArticle A again")
            st.caption("Pasting keeps cell text, column order and blank cells. Clipboard paste does not include Excel colours, fonts or widths.")
            if not text:
                return
            grid = parse_paste(text)
            identity = text.encode()
        else:
            upload = st.file_uploader("Choose an Excel workbook", type=["xlsx"])
            st.caption("For plain worksheets, upload keeps cell formatting and other sheets. Advanced workbooks may need to be pasted as plain data.")
            if upload is None:
                return
            source = upload.getvalue()
            wb = load_workbook(BytesIO(source))
            sheet = st.selectbox("Worksheet to clean", wb.sheetnames)
            ws = wb[sheet]
            if ws.max_row * ws.max_column > 500_000:
                st.error("This sheet is too large for interactive review. Use a smaller range.")
                return
            grid = [[c.value for c in row] for row in ws]
            identity = source
        c1, c2 = st.columns(2)
        layout = c1.selectbox("Where are the URLs?", ["Down a column (delete duplicate rows)",
                                                     "Across a row (delete duplicate columns)"])
        axis = "Rows" if layout.startswith("Down") else "Columns"
        limit = len(grid[0]) if axis == "Rows" else len(grid)
        position = c2.number_input("URL column number" if axis == "Rows" else "URL row number",
                                   min_value=1, max_value=limit, value=1) - 1
        a, b = st.columns(2)
        skip = a.number_input("Header rows to keep" if axis == "Rows" else "Leading columns to keep",
                              min_value=0, max_value=len(grid) if axis == "Rows" else len(grid[0]),
                              value=1 if axis == "Rows" else 0)
        trim = b.checkbox("Ignore spaces before / after URLs", value=True)
        st.caption("Matching is otherwise exact: case, http/https, trailing slashes and query parameters remain distinct. Blank URLs are kept.")
        token = hashlib.sha256(identity + repr((sheet, axis, position, skip, trim)).encode()).hexdigest()
        if st.session_state.get("input_token") != token:
            st.session_state.input_token = token
            st.session_state.removed = set()
            st.session_state.revision = 0
            st.session_state.select_all = True
        removed = st.session_state.removed
        records = grid if axis == "Rows" else list(map(list, zip(*grid)))
        # Recheck remaining records; original indices remain stable after deletions.
        live_ids = [i for i in range(len(records)) if i not in removed]
        live_grid = [grid[i] for i in live_ids] if axis == "Rows" else [[r[i] for i in live_ids] for r in grid]
        live_skip = sum(i < skip for i in live_ids)
        _, _, status = inspect_duplicates(live_grid, axis, position, live_skip, trim)
        status_map = dict(zip(live_ids, status))
        duplicates = [i for i in live_ids if status_map[i] == "Duplicate"]
        m1, m2, m3 = st.columns(3)
        m1.metric(f"Remaining {axis.lower()}", len(live_ids))
        m2.metric("Repeated entries", len(duplicates))
        m3.metric("Deleted", len(removed))
        labels = [f"Column {get_column_letter(i + 1)}" for i in range(len(grid[0]))]
        preview = pd.DataFrame([["" if v is None else str(v) for v in r] for r in live_grid], columns=labels)
        preview.index = [i + 1 for i in live_ids] if axis == "Rows" else range(1, len(grid) + 1)

        def highlight(data):
            styles = pd.DataFrame("", index=data.index, columns=data.columns)
            for i in live_ids:
                color = "#ffe0df" if status_map[i] == "Duplicate" else "#fff0bc" if status_map[i] == "First occurrence" else ""
                if color:
                    if axis == "Rows":
                        styles.loc[i + 1, :] = f"background-color: {color}; color: #222"
                    else:
                        styles.iloc[:, live_ids.index(i)] = f"background-color: {color}; color: #222"
            return styles

        st.subheader("Review your data")
        st.caption("Yellow = first occurrence to keep. Red = repeated occurrence. Numbers refer to the original input.")
        st.dataframe(preview.style.apply(highlight, axis=None), use_container_width=True)
        grouped_ids = [i for i in live_ids if status_map[i] in ("Duplicate", "First occurrence")]
        if grouped_ids:
            st.subheader("Choose entries to delete")
            st.caption("The first occurrence is kept. Tick the repeated entries you want to remove; the complete row or column is deleted.")
            left, right, reset = st.columns(3)
            if left.button("Select all repeated entries"):
                st.session_state.select_all = True
                st.session_state.revision += 1
                st.rerun()
            if right.button("Clear selection"):
                st.session_state.select_all = False
                st.session_state.revision += 1
                st.rerun()
            if reset.button("Restore all deleted entries"):
                st.session_state.removed = set()
                st.session_state.revision += 1
                st.rerun()
            choices = pd.DataFrame({"Delete": [st.session_state.select_all] * len(duplicates),
                                    "Original position": [i + 1 for i in duplicates],
                                    "URL": [str(records[i][position]) for i in duplicates],
                                    "Other cells": [" | ".join(str(v) if v is not None else "" for j, v in enumerate(records[i]) if j != position) for i in duplicates]})
            edited = st.data_editor(choices, hide_index=True, use_container_width=True,
                                    disabled=["Original position", "URL", "Other cells"],
                                    column_config={"Delete": st.column_config.CheckboxColumn("Delete")},
                                    key=f"review_{token}_{st.session_state.revision}")
            selected = {int(v) - 1 for v in edited.loc[edited["Delete"], "Original position"]}
            if st.button(f"Delete {len(selected)} selected {axis.lower()}", disabled=not selected, type="primary"):
                removed.update(selected)
                st.session_state.revision += 1
                st.rerun()
        elif removed and st.button("Restore all deleted entries"):
            st.session_state.removed = set()
            st.session_state.revision += 1
            st.rerun()
        if duplicates:
            st.info("Delete all remaining repeated entries to enable the non-duplicate Excel download.")
        else:
            st.success("No duplicate URLs remain. Your Excel is ready.")
            try:
                output = export_excel(grid, removed, axis, source, sheet)
                st.download_button("Download cleaned Excel", output, "cleaned_urls.xlsx",
                                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            except ValueError as exc:
                st.error(str(exc))
    except Exception as exc:
        st.error(f"Could not read the data: {exc}")


def main():
    st.set_page_config(page_title="URL Tools", page_icon="🔗", layout="wide")
    page = st.sidebar.radio("Page", ["Duplicate Cleaner", "URL Analyzer"])
    if page == "Duplicate Cleaner":
        cleaner_page()
    else:
        analyzer_page()


if __name__ == "__main__":
    main()
